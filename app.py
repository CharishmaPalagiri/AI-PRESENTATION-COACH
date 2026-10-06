import re
import requests
import pandas as pd
import streamlit as st
from pptx import Presentation
import speech_recognition as sr

from auth import register_user, login_user, save_session, get_sessions

# ---------------- SETTINGS ----------------
OLLAMA_URL = "http://localhost:11434/api/generate"
OLLAMA_MODEL = "llama3.2:3b"
FILLERS = ["um", "uh", "like", "you know", "basically", "actually", "kind of", "sort of"]

st.set_page_config(page_title="AI Presentation Coach", page_icon="🎤")

if "user" not in st.session_state:
    st.session_state.user = None


# ---------------- COACH FUNCTIONS ----------------
def analyze(text, wpm=130):
    words = re.findall(r"\b\w+\b", text.lower())
    sentences = [s for s in re.split(r"[.!?]+", text) if s.strip()]
    filler_count = sum(len(re.findall(rf"\b{f}\b", text.lower())) for f in FILLERS)
    return {
        "words": len(words),
        "sentences": len(sentences),
        "avg_sentence": round(len(words) / max(len(sentences), 1), 1),
        "fillers": filler_count,
        "minutes": round(len(words) / wpm, 2),
    }


def ask_ollama(text, stats):
    prompt = f"""You are a friendly presentation coach for students.
Speech/slide text:
{text}

Stats: {stats}

Give feedback in this format:
1. Strengths
2. Weaknesses
3. Filler word advice
4. Three tips to improve
5. Score out of 10
Keep it short and simple."""
    try:
        r = requests.post(
            OLLAMA_URL,
            json={"model": OLLAMA_MODEL, "prompt": prompt, "stream": False},
            timeout=300,
        )
        return r.json()["response"]
    except Exception as e:
        return f"Could not reach Ollama. Is it running? Error: {e}"


def extract_pptx(file):
    prs = Presentation(file)
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


def transcribe_wav(file):
    r = sr.Recognizer()
    with sr.AudioFile(file) as source:
        audio = r.record(source)
    return r.recognize_google(audio)  # needs internet


# ---------------- REGISTER + LOGIN PAGE ----------------
def login_page():
    st.title("🎤 AI Presentation Coach")

    # Register is the FIRST tab, Login is the SECOND tab
    tab_register, tab_login = st.tabs(["Register", "Login"])

    with tab_register:
        with st.form("register_form"):
            u = st.text_input("Choose username")
            p = st.text_input("Choose password", type="password")
            if st.form_submit_button("Register"):
                if not u.strip() or not p:
                    st.warning("Please fill both fields")
                elif register_user(u, p):
                    st.success("Account created! Now click the Login tab.")
                else:
                    st.error("Username already exists")

    with tab_login:
        with st.form("login_form"):
            u = st.text_input("Username")
            p = st.text_input("Password", type="password")
            if st.form_submit_button("Login"):
                if login_user(u, p):
                    st.session_state.user = u.strip().lower()
                    st.rerun()
                else:
                    st.error("Wrong username or password")


# ---------------- COACH PAGE ----------------
def coach_page():
    st.header("Practice your presentation")
    mode = st.radio(
        "Input type",
        ["Type / paste text", "Upload PPTX", "Upload WAV audio"],
        horizontal=True,
    )
    text = ""

    if mode == "Type / paste text":
        text = st.text_area("Paste your speech here", height=200)
    elif mode == "Upload PPTX":
        f = st.file_uploader("Choose a .pptx file", type=["pptx"])
        if f:
            text = extract_pptx(f)
            st.text_area("Extracted text", text, height=200)
    else:
        f = st.file_uploader("Choose a .wav file", type=["wav"])
        if f:
            with st.spinner("Converting speech to text..."):
                try:
                    text = transcribe_wav(f)
                    st.text_area("Transcript", text, height=200)
                except Exception as e:
                    st.error(f"Could not transcribe: {e}")

    title = st.text_input("Session title", "My presentation")

    if st.button("Analyze with AI", type="primary"):
        if not text.strip():
            st.warning("Please give some text first")
            return

        stats = analyze(text)
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Words", stats["words"])
        c2.metric("Minutes", stats["minutes"])
        c3.metric("Filler words", stats["fillers"])
        c4.metric("Avg sentence", stats["avg_sentence"])

        with st.spinner("Coach is thinking (Ollama)..."):
            feedback = ask_ollama(text, stats)

        st.subheader("AI Feedback")
        st.write(feedback)
        save_session(st.session_state.user, title, text, stats, feedback)
        st.success("Saved to history")


# ---------------- HISTORY PAGE ----------------
def history_page():
    st.header("Your past sessions")
    data = get_sessions(st.session_state.user)
    if not data:
        st.info("No sessions yet.")
        return

    df = pd.DataFrame(data)[["created", "title", "words", "minutes", "fillers"]]
    st.dataframe(df)
    for d in data:
        with st.expander(f"{d['created']} - {d['title']}"):
            st.write(d["feedback"])


# ---------------- MAIN ----------------
if st.session_state.user is None:
    login_page()
else:
    st.sidebar.success(f"Logged in as {st.session_state.user}")
    page = st.sidebar.radio("Menu", ["Coach", "History"])
    if st.sidebar.button("Logout"):
        st.session_state.user = None
        st.rerun()

    if page == "Coach":
        coach_page()
    else:
        history_page()